#!/usr/bin/env python3
"""xlsxの数式を実際に計算させて、エラーが出ないか見る。

このセッションのサンドボックスではLibreOfficeが動かない（3セルのファイルでも
固まる）ので、検算はGitHub Actionsのランナーに回す。
openpyxlは数式を文字列で書くだけで計算結果を持たないため、
一度LibreOfficeに通さないと「本当に計算できる式か」が分からない。
"""
import subprocess
import sys
import pathlib
import shutil

ERRORS = ("#REF!", "#VALUE!", "#NAME?", "#DIV/0!", "#N/A", "#NULL!", "#NUM!", "Err:")


def main(path):
    src = pathlib.Path(path).resolve()
    out = src.parent / "_recalc"
    if out.exists():
        shutil.rmtree(out)
    out.mkdir()

    # LibreOfficeは読み込み時に再計算する。同じ形式で書き出せば結果が入る
    subprocess.run(
        ["soffice", "--headless", "--norestore", "--nolockcheck",
         "--convert-to", "xlsx", "--outdir", str(out), str(src)],
        check=True, timeout=600,
    )

    done = out / src.name
    if not done.exists():
        raise SystemExit(f"！ 変換されなかった: {done}")

    from openpyxl import load_workbook
    wb = load_workbook(done, data_only=True)

    総数 = 空 = 0
    見つけたエラー = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                総数 += 1
                v = str(c.value)
                if any(e in v for e in ERRORS):
                    見つけたエラー.append(f"{ws.title}!{c.coordinate} = {v}")

    # 数式が本当に計算されたかを抜き取りで見る
    抜き取り = [
        ("日次入力", "E5", "客単価（売上÷来客数）"),
        ("日次入力", "B5", "曜日"),
        ("日次入力", "C3", "売上合計"),
        ("月次集計", "A7", "1か月目の月"),
        ("月次集計", "B7", "1か月目の売上合計"),
        ("商品・サービス別", "D5", "粗利"),
        ("商品・サービス別", "E5", "粗利率"),
    ]
    print("値が入っているセル:", 総数)
    print()
    for sheet, cell, label in 抜き取り:
        val = wb[sheet][cell].value
        印 = "○" if val not in (None, "") else "！空"
        print(f"  {印} {sheet}!{cell:5} {label:28} = {val!r}")

    print()
    if 見つけたエラー:
        print(f"！ エラーのセル {len(見つけたエラー)}件")
        for e in 見つけたエラー[:50]:
            print("   ", e)
        raise SystemExit(1)
    print("○ エラーのセルは無し")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "assets/downloads/uriage-kanri-template.xlsx")
